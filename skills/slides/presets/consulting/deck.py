#!/usr/bin/env python3
"""Consulting preset: a case or strategy deck at 13.333 x 7.5 in, built with python-pptx.

Copy this folder next to the deck's images, replace the demo slides under SLIDES, and run:

    uv run --no-project --with python-pptx --with pillow python deck.py out.pptx
    SLIDES=1,4 ... python deck.py out.pptx     only those slides (1-based, in deck order)

Every text box is checked against the real Outfit metrics: a box whose text would overflow, or
a line marked maxlines=1 that would wrap, prints a WARN line. Treat each one as a bug. The
look and the writing rules are in STYLE.md next to this file.

The demo deck argues a fictional case (Fernwood, a coffee-pod brand) so that every pattern
appears once with realistic text. Its company, people and numbers are invented."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.ns import qn
import os, re, shutil, sys, zipfile

try:
    from PIL import ImageFont
except ImportError:
    ImageFont = None

# ---------------------------------------------------------------- palette
INK   = RGBColor(0x10,0x18,0x24)   # body text
INKM  = RGBColor(0x2E,0x38,0x44)   # secondary body text
MUTE  = RGBColor(0x4F,0x5B,0x66)   # subtitles, column heads
FAINT = RGBColor(0x66,0x72,0x7E)   # source line
RULE  = RGBColor(0x87,0x95,0xA0)   # row rules (dark enough for a projector)
RULE2 = RGBColor(0x5E,0x6B,0x76)   # header rules, arrows
NAVY  = RGBColor(0x0B,0x14,0x33)   # title and recommendation slides
PRIM  = RGBColor(0x22,0x3A,0xE6)   # the one accent: labels, banners, bars
PRIMD = RGBColor(0x16,0x25,0x9E)
PTINT = RGBColor(0xE6,0xE9,0xFC)   # highlighted row, light stage box
WARN  = RGBColor(0xD2,0x45,0x15)   # the problem, the critic, "not delivered"
WARNT = RGBColor(0xFB,0xE9,0xE0)
GOOD  = RGBColor(0x15,0x70,0x4A)
GREY  = RGBColor(0xC9,0xD1,0xD8)
CARD  = RGBColor(0xF2,0xF4,0xF7)
WHITE = RGBColor(0xFF,0xFF,0xFF)
SKYL  = RGBColor(0xD3,0xDC,0xF6)   # body text on navy
SKYD  = RGBColor(0xA9,0xB9,0xEE)   # labels on navy
NAVYL = RGBColor(0x3A,0x4B,0x8A)   # rules on navy
REG="Outfit"; BODYF="Outfit Medium"   # static weights; variable Outfit renders thin
W,H=13.333,7.5
ML=0.75; CW=W-2*ML
TOP=MSO_ANCHOR.TOP; MID=MSO_ANCHOR.MIDDLE; BOT=MSO_ANCHOR.BOTTOM
LEFT=PP_ALIGN.LEFT; CENTER=PP_ALIGN.CENTER

_D=os.path.dirname(os.path.abspath(__file__))
_FONT_DIRS=[os.environ.get("FONT_DIR",""),os.path.expanduser("~/Library/Fonts"),"/Library/Fonts","/System/Library/Fonts/Supplemental",
            os.path.expanduser("~/.local/share/fonts"),os.path.expanduser("~/.fonts"),"/usr/share/fonts/truetype/outfit"]
_FF={}
WARNINGS=[]
CUR={"n":0,"name":""}

def _font_file(name):
    for d in _FONT_DIRS:
        if d and os.path.exists(os.path.join(d,name)): return os.path.join(d,name)
    return None

def _measure_font(name,bold,size):
    if name.startswith("Outfit"):
        files=["Outfit-Bold.ttf" if bold else {"Outfit":"Outfit-Regular.ttf","Outfit Medium":"Outfit-Medium.ttf"}.get(name,"Outfit-Regular.ttf")]
    else:  # a template's font: try the usual file names
        files=[f"{name} Bold.ttf",f"{name}-Bold.ttf",f"{name}Bold.ttf"] if bold else [f"{name}.ttf",f"{name}-Regular.ttf",f"{name} Regular.ttf"]
    k=(files[0],size)
    if k not in _FF:
        path=next((x for x in map(_font_file,files) if x),None)
        if ImageFont is None or path is None:
            _FF[k]=None
            msg=f"fit checks off: no font file for {name}; install the static Outfit weights (or the template's font) or set FONT_DIR"
            if msg not in WARNINGS: WARNINGS.append(msg)
        else:
            _FF[k]=ImageFont.truetype(path,size=max(1,int(round(size*10))))
    return _FF[k]

# ---------------------------------------------------------------- template mode
# TEMPLATE=company.potx builds on the template itself: its masters, logos and footers stay, slides go on its
# emptiest layout, the theme's fonts and colors replace the preset's, and every slide is scaled uniformly from the
# 13.333 x 7.5 canvas below to the template's size (top-aligned, so a 4:3 template keeps its footer band free).
# Run ../../scripts/template.py on it first to see where its logo and footer sit.
TEMPLATE=os.environ.get("TEMPLATE")
SCALE=1.0
def _hex(c): return RGBColor.from_string(c)
def _mix(c,t,a): return RGBColor(*[round(int(str(c)[i:i+2],16)*(1-a)+int(str(t)[i:i+2],16)*a) for i in (0,2,4)])
if TEMPLATE:
    for _p in (os.path.join(_D,"..","..","scripts"),os.path.expanduser("~/.claude/skills/slides/scripts")):
        sys.path.append(_p)  # the skill's scripts/, whether this file sits in the skill or was copied next to a deck
    from template import open_template, theme, blank_layout, lum
    prs=open_template(TEMPLATE)
    for sid in list(prs.slides._sldIdLst):  # drop the template's own slides
        prs.part.drop_rel(sid.rId); prs.slides._sldIdLst.remove(sid)
    BLANK=blank_layout(prs)
    _fonts,_cols=theme(prs)
    REG=BODYF=_fonts.get("minor") or REG
    if _cols.get("dk1"): INK=INKM=_hex(_cols["dk1"])
    if _cols.get("accent1"):
        PRIM=_hex(_cols["accent1"]); PRIMD=_mix(PRIM,RGBColor(0,0,0),0.3); PTINT=_mix(PRIM,WHITE,0.88)
    if _cols.get("dk1") and lum(_cols["dk1"])<0.2: NAVY=_hex(_cols["dk1"])
    SCALE=min(prs.slide_width/Inches(W),prs.slide_height/Inches(H))
else:
    prs=Presentation(); prs.slide_width=Inches(W); prs.slide_height=Inches(H)
    BLANK=prs.slide_layouts[6]

def _rescale(sl,f):
    """Scale every shape, font and line on one slide by f (template mode)."""
    for sh in sl.shapes:
        for a in ("left","top","width","height"):
            v=getattr(sh,a)
            if v is not None: setattr(sh,a,int(v*f))
        for r in sh._element.iter(qn("a:rPr"),qn("a:defRPr")):
            if r.get("sz"): r.set("sz",str(max(100,int(int(r.get("sz"))*f))))
        for ln in sh._element.iter(qn("a:ln")):
            if ln.get("w"): ln.set("w",str(int(int(ln.get("w"))*f)))
        for sp in sh._element.iter(qn("a:spcPts")):
            if sp.get("val"): sp.set("val",str(int(int(sp.get("val"))*f)))

# ---------------------------------------------------------------- primitives
def slide(bg=WHITE):
    s=prs.slides.add_slide(BLANK)
    for ph in list(s.placeholders): ph._element.getparent().remove(ph._element)  # footer/date placeholders
    if not (TEMPLATE and bg==WHITE):  # in template mode a white slide shows the master's own background
        s.background.fill.solid(); s.background.fill.fore_color.rgb=bg
    return s
def hide(s): s._element.set("show","0")     # backup slides: in the file, skipped when presenting
def notes(s,txt): s.notes_slide.notes_text_frame.text=txt

def _font(run,size,color,bold,font=BODYF,italic=False):
    tf=REG if bold else font
    run.font.size=Pt(size); run.font.name=tf; run.font.bold=bold; run.font.italic=italic
    run.font.color.rgb=color
    rPr=run._r.get_or_add_rPr()
    for tag in ("a:latin","a:cs"):
        e=rPr.find(qn(tag))
        if e is None: e=rPr.makeelement(qn(tag),{}); rPr.append(e)
        e.set("typeface",tf)

def _bullet(p,color,indent=0.34):
    pPr=p._p.get_or_add_pPr()
    pPr.set("marL",str(int(Inches(indent)))); pPr.set("indent",str(-int(Inches(indent))))
    for tag in ("a:buClr","a:buSzPct","a:buFont","a:buChar","a:buNone"):
        e=pPr.find(qn(tag))
        if e is not None: pPr.remove(e)
    clr=pPr.makeelement(qn("a:buClr"),{}); srgb=clr.makeelement(qn("a:srgbClr"),{"val":str(color)}); clr.append(srgb)
    pPr.append(clr)
    pPr.append(pPr.makeelement(qn("a:buSzPct"),{"val":"80000"}))
    pPr.append(pPr.makeelement(qn("a:buFont"),{"typeface":"Arial"}))
    pPr.append(pPr.makeelement(qn("a:buChar"),{"char":"●"}))

def _fill_text(tf,runs,size,color,bold,align,font,sp,italic):
    """runs: str | list of (text, overrides) | "\n" (new paragraph). Overrides: size, color, bold,
    font, italic, newp, align, spc_before, bullet, bullet_color."""
    if isinstance(runs,str): runs=[(runs,{})]
    paras=[]; first=True; p=None
    for item in runs:
        if item=="\n":
            p=tf.add_paragraph(); p.alignment=align; paras.append([p,[],{}]); continue
        t,ov=item
        if first:
            p=tf.paragraphs[0]; p.alignment=ov.get("align",align); first=False; paras.append([p,[],ov])
        elif ov.get("newp"):
            p=tf.add_paragraph(); p.alignment=ov.get("align",align); paras.append([p,[],ov])
        if sp!=1.0: p.line_spacing=sp
        if ov.get("spc_before") is not None: p.space_before=Pt(ov["spc_before"])
        if ov.get("bullet"): _bullet(p,ov.get("bullet_color",ov.get("color",color)))
        r=p.add_run(); r.text=t
        sz=ov.get("size",size); b=ov.get("bold",bold); fn=ov.get("font",font)
        _font(r,sz,ov.get("color",color),b,fn,ov.get("italic",italic))
        paras[-1][1].append((t,sz,b,fn))
    return paras

def _estimate(paras,w_in,sp,indent_bullets=0.34):
    """Greedy wrap with the real Outfit metrics; returns (height_pt, max_lines_per_para) or None.
    Google Slides sets Outfit about 10% wider than these metrics, so leave slack."""
    total=0.0; maxlines=0
    for p,runs,ov in paras:
        if not runs: total+=12*1.2*sp; continue
        width=(w_in-(indent_bullets if ov.get("bullet") else 0))*72
        words=[]
        for t,sz,b,fn in runs:
            f=_measure_font(fn,b,sz)
            if f is None: return None
            for wd in re.split(r"(\s+)",t):
                if wd=="": continue
                words.append((wd,f.getlength(wd)/10.0,sz))
        lines=1; cur=0.0; msz=max(sz for _,sz,_,_ in runs)
        for wd,wl,sz in words:
            if wd.isspace(): cur+=wl; continue
            if cur>0 and cur+wl>width+0.5: lines+=1; cur=wl
            else: cur+=wl
        total+=lines*msz*1.2*sp+(ov.get("spc_before") or 0)
        maxlines=max(maxlines,lines)
    return total,maxlines

def text(s,x,y,w,h,runs,size=22,color=INK,bold=False,align=LEFT,anchor=TOP,font=BODYF,sp=1.0,
         italic=False,fill=None,line=None,lw=1.0,margin=0.0,shape=None,maxlines=None):
    """One text box that carries its own fill and border; never a rectangle under a separate box."""
    if shape is None:
        tb=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    else:
        tb=s.shapes.add_shape(shape,Inches(x),Inches(y),Inches(w),Inches(h)); tb.shadow.inherit=False
    if fill is not None: tb.fill.solid(); tb.fill.fore_color.rgb=fill
    elif shape is not None: tb.fill.background()
    if line is not None: tb.line.color.rgb=line; tb.line.width=Pt(lw)
    elif shape is not None: tb.line.fill.background()
    tf=tb.text_frame; tf.word_wrap=True; tf.vertical_anchor=anchor; tf.auto_size=MSO_AUTO_SIZE.NONE
    ml_,mr_,mt_,mb_=(margin,)*4 if isinstance(margin,(int,float)) else margin
    tf.margin_left=Inches(ml_); tf.margin_right=Inches(mr_); tf.margin_top=Inches(mt_); tf.margin_bottom=Inches(mb_)
    paras=_fill_text(tf,runs,size,color,bold,align,font,sp,italic)
    est=_estimate(paras,w-ml_-mr_,sp)
    if est:
        hgt,ml=est
        if hgt>(h-mt_-mb_)*72+2:
            WARNINGS.append(f"slide {CUR['n']} {CUR['name']}: overflow {hgt/72:.2f}in > {h-mt_-mb_:.2f}in :: {str(runs)[:70]}")
        if maxlines is not None and ml>maxlines:
            WARNINGS.append(f"slide {CUR['n']} {CUR['name']}: wraps to {ml} lines (max {maxlines}) :: {str(runs)[:70]}")
    return tb

def rect(s,x,y,w,h,fill=None,line=None,lw=1.0,shape=MSO_SHAPE.RECTANGLE,dash=False):
    sp=s.shapes.add_shape(shape,Inches(x),Inches(y),Inches(w),Inches(h)); sp.shadow.inherit=False
    if fill is not None: sp.fill.solid(); sp.fill.fore_color.rgb=fill
    else: sp.fill.background()
    if line is not None:
        sp.line.color.rgb=line; sp.line.width=Pt(lw)
        if dash: sp.line.dash_style=MSO_LINE_DASH_STYLE.DASH
    else: sp.line.fill.background()
    return sp
def hline(s,x1,y,x2,color=RULE,lw=1.25):
    ln=s.shapes.add_connector(1,Inches(x1),Inches(y),Inches(x2),Inches(y))
    ln.line.color.rgb=color; ln.line.width=Pt(lw); ln.shadow.inherit=False; return ln
def dot(s,cx,cy,r,fill):
    return rect(s,cx-r,cy-r,2*r,2*r,fill=fill,shape=MSO_SHAPE.OVAL)
def ring(s,cx,cy,r,color,lw,fill=None,dash=False):
    return rect(s,cx-r,cy-r,2*r,2*r,fill=fill,line=color,lw=lw,shape=MSO_SHAPE.OVAL,dash=dash)
def picture(s,file,x,y,w,ratio):
    """An image at width w (ratio = height / width), with a thin rule around it."""
    s.shapes.add_picture(os.path.join(_D,file),Inches(x),Inches(y),Inches(w),Inches(w*ratio))
    rect(s,x,y,w,w*ratio,line=RULE,lw=1.0)

# ---------------------------------------------------------------- slide furniture
def head(s,title,sub=None,dark=False):
    """Short title and an italic subtitle in ONE text box, then a rule with a short accent."""
    tc=WHITE if dark else INK; sc=SKYD if dark else MUTE
    runs=[(title,{"size":38,"color":tc,"font":BODYF})]
    if sub: runs.append((sub,{"newp":True,"size":24,"italic":True,"color":sc,"font":REG,"spc_before":4}))
    text(s,ML,0.36,CW,1.3,runs,anchor=TOP,maxlines=1)
    hline(s,ML,1.74,W-ML,(NAVYL if dark else RULE),1.5)
    rect(s,ML,1.715,1.3,0.055,fill=PRIM)
def src(s,t,dark=False):
    """Where the slide's evidence comes from, one line at the foot. No page numbers here."""
    text(s,ML,7.02,CW,0.3,t,size=15,color=(SKYD if dark else FAINT),font=REG,maxlines=1)
def banner(s,t,y=6.2,size=22):
    """The takeaway in a solid accent bar. Use it where the slide has one sentence to land, not on every slide."""
    text(s,ML,y,CW,0.66,t,size=size,color=WHITE,bold=True,fill=PRIM,anchor=MID,align=CENTER,
         margin=(0.2,0.2,0.05,0.05),maxlines=1)

# ---------------------------------------------------------------- patterns
def numbered_rows(s,rows,y=2.02,rh=1.02):
    """01 / name / what it means: a standard of comparison, a set of tests, the agenda."""
    for i,(t,d) in enumerate(rows):
        cy=y+i*rh
        text(s,ML,cy,0.85,rh,f"0{i+1}",size=30,color=PRIM,bold=True,anchor=MID,maxlines=1)
        text(s,ML+0.95,cy,3.0,rh,t,size=26,color=INK,bold=True,anchor=MID,maxlines=1)
        text(s,ML+4.05,cy,CW-4.05,rh,d,size=24,color=INKM,anchor=MID,maxlines=1)
        if i<len(rows)-1: hline(s,ML,cy+rh,W-ML,RULE,1.25)

def label_rows(s,rows,y=2.98,rh=0.62,lw=3.1,size=21):
    """Label: evidence rows. The independent legs of one argument, each from its own source."""
    for i,(lab,t) in enumerate(rows):
        cy=y+i*rh
        text(s,ML,cy,lw,rh,lab,size=size,color=PRIM,bold=True,anchor=MID,maxlines=1)
        text(s,ML+lw+0.1,cy,CW-lw-0.1,rh,t,size=size,color=INK,anchor=MID,maxlines=1)
        if i<len(rows)-1: hline(s,ML,cy+rh,W-ML,RULE,1.0)

def rule_table(s,heads,rows,widths,y=1.98,rh=0.72,size=20,styles=None):
    """A table drawn with rules, not cell borders. rows: tuples of cells; a cell is text or
    (text, color). styles: per column (color, bold), default first column accent + bold."""
    styles=styles or [(PRIM,True)]+[(INK,False)]*(len(widths)-1)
    xs=[ML+sum(widths[:j]) for j in range(len(widths))]
    for hd,x,w in zip(heads,xs,widths):
        if hd: text(s,x,y,w-0.15,0.4,hd,size=18,color=MUTE,bold=True,maxlines=1)
    hline(s,ML,y+0.46,W-ML,RULE2,1.5)
    y2=y+0.56
    for i,row in enumerate(rows):
        cy=y2+i*rh
        for j,(cell,x,w) in enumerate(zip(row,xs,widths)):
            col,b=styles[j]
            if isinstance(cell,tuple): cell,col=cell; b=True
            text(s,x,cy,w-0.2,rh,cell,size=size,color=col,bold=b,anchor=MID,sp=1.0)
        if i<len(rows)-1: hline(s,ML,cy+rh,W-ML,RULE,1.0)
    return y2+len(rows)*rh

def share_bar(s,y,segs,h=0.72):
    """One 100% bar: [(pct, label, fill, text color)]. Highlight the segment the argument needs."""
    x=ML
    for pct,lab,fill,tc in segs:
        w=CW*pct/100.0
        text(s,x,y,w,h,f"{pct}%",size=22,color=tc,bold=True,align=CENTER,anchor=MID,fill=fill,line=WHITE,lw=1.5,
             shape=MSO_SHAPE.RECTANGLE,margin=0.02,maxlines=1)
        text(s,x,y+h+0.06,w,0.36,lab,size=17,color=MUTE,align=CENTER,maxlines=1)
        x+=w

def stages(s,items,y=2.05):
    """A left-to-right process (a customer decision process, a pipeline). items: (name, body, emphasis)
    where emphasis is None, "accent" or "warn"."""
    n=len(items); gap=0.26; cw=(CW-(n-1)*gap)/n
    for i,(st,body,emph) in enumerate(items):
        x=ML+i*(cw+gap)
        fill,tc,bc={None:(PTINT,PRIM,INK),"accent":(PRIM,WHITE,PRIM),"warn":(WARN,WHITE,WARN)}[emph]
        text(s,x,y,cw,0.95,st,size=19,color=tc,bold=True,align=CENTER,anchor=MID,fill=fill,margin=(0.08,0.08,0.04,0.04),sp=0.95)
        if i<n-1: text(s,x+cw,y,gap,0.95,"›",size=26,color=RULE2,bold=True,align=CENTER,anchor=MID)
        text(s,x,y+1.12,cw,1.85,body,size=20,color=bc,bold=emph is not None,align=CENTER,sp=1.0)
    return [ML+i*(cw+gap) for i in range(n)],cw

def scorecard(s,heads,rows,y=1.95,nw=6.4,rh=0.56):
    """Candidates scored 0-3 on each criterion with dots; rows: (name, [scores], is_ours)."""
    cw=(CW-nw)/len(heads)
    text(s,ML,y,nw,0.62,"Option",size=18,color=MUTE,bold=True,anchor=BOT,maxlines=1)
    for j,hd in enumerate(heads):
        text(s,ML+nw+j*cw,y,cw,0.62,hd,size=18,color=MUTE,bold=True,align=CENTER,anchor=BOT,maxlines=2,sp=0.95)
    hline(s,ML,y+0.7,W-ML,RULE2,1.5)
    for i,(nm,scores,win) in enumerate(rows):
        cy=y+0.78+i*rh
        if win: rect(s,ML-0.1,cy+0.02,CW+0.2,rh-0.04,fill=PTINT)
        text(s,ML,cy,nw-0.15,rh,nm,size=19,color=INK,bold=win,anchor=MID,maxlines=1)
        for j,lv in enumerate(scores):
            bx=ML+nw+j*cw+cw/2-0.27
            for k in range(3): dot(s,bx+k*0.27,cy+rh/2,0.09,PRIM if k<lv else GREY)
        if i<len(rows)-1: hline(s,ML,cy+rh,W-ML,RULE,0.75)

def objection_rows(s,pairs,y=2.0,rh=0.92,c0=5.15):
    """Their best answer → why the point holds."""
    for i,(q,a) in enumerate(pairs):
        cy=y+i*rh
        text(s,ML,cy,c0,rh,q,size=21,color=WARN,bold=True,anchor=MID,sp=1.0)
        text(s,ML+c0+0.05,cy,0.5,rh,"→",size=22,color=RULE2,bold=True,align=CENTER,anchor=MID,maxlines=1)
        text(s,ML+c0+0.65,cy,CW-c0-0.65,rh,a,size=21,color=INK,anchor=MID,sp=1.0)
        if i<len(pairs)-1: hline(s,ML,cy+rh,W-ML,RULE,1.0)

def horizons(s,cols,y=2.0):
    """Short / medium / long term on a navy slide: (horizon, headline, bullets, why, owner, metric)."""
    gap=0.45; cw=(CW-2*gap)/3
    for i,(hz,hd,bul,why,owner,metric) in enumerate(cols):
        x=ML+i*(cw+gap)
        text(s,x,y,cw,0.4,hz,size=18,color=SKYD,bold=True,maxlines=1)
        hline(s,x,y+0.46,x+cw,NAVYL,1.25)
        text(s,x,y+0.56,cw,0.52,hd,size=24,color=WHITE,bold=True,maxlines=1)
        runs=[(b,{"size":20,"color":SKYL,"bullet":True,"bullet_color":SKYD,"newp":k>0,"spc_before":(0 if k==0 else 4)}) for k,b in enumerate(bul)]
        text(s,x,y+1.15,cw,1.7,runs,sp=1.0)
        text(s,x,y+2.88,cw,0.8,[("Why: ",{"size":19,"color":SKYD,"bold":True}),(why,{"size":19,"color":WHITE})],sp=1.0)
        text(s,x,y+3.74,cw,0.95,[("Owner: ",{"size":18,"color":SKYD,"bold":True}),(owner,{"size":18,"color":WHITE,"bold":True}),
            ("Metric: ",{"size":18,"color":SKYD,"bold":True,"newp":True,"spc_before":3}),(metric,{"size":18,"color":WHITE,"bold":True})],sp=1.0)

def big_numbers(s,stats,y=5.72):
    """Two to four large figures with what each counts underneath."""
    cw=CW/len(stats)
    for i,(n,lab) in enumerate(stats):
        x=ML+i*cw
        text(s,x,y,cw-0.2,0.75,n,size=40,color=PRIM,bold=True,anchor=MID,maxlines=1)
        text(s,x,y+0.72,cw-0.2,0.38,lab,size=17,color=MUTE,maxlines=1)

def flow(s,nodes,edges,size=19):
    """A flowchart whose arrows are connectors glued to the boxes, so they follow a box moved in
    Google Slides. nodes: {key: (x, y, w, h, label, style)} with style None | "accent" | "warn" | "ghost";
    edges: (from, to, label) or (from, to, label, (side_from, side_to)); sides are 0 top, 1 left, 2 bottom,
    3 right, picked from the relative position when not given. Straight between facing sides, elbow otherwise."""
    look={None:(PTINT,INK,None),"accent":(PRIM,WHITE,None),"warn":(WARNT,WARN,WARN),"ghost":(WHITE,MUTE,RULE)}
    shp={}
    for k,(x,y,w,h,label,st) in nodes.items():
        fill,tc,ln=look[st]
        shp[k]=text(s,x,y,w,h,label,size=size,color=tc,bold=True,align=CENTER,anchor=MID,fill=fill,line=ln,lw=1.25,
                    shape=MSO_SHAPE.RECTANGLE,margin=(0.1,0.1,0.05,0.05),sp=0.95)
    for e in edges:
        a,b,label=e[:3]
        A,B=nodes[a],nodes[b]
        dx=(B[0]+B[2]/2)-(A[0]+A[2]/2); dy=(B[1]+B[3]/2)-(A[1]+A[3]/2)
        sa,sb=e[3] if len(e)>3 else (((3,1) if dx>0 else (1,3)) if abs(dx)>=abs(dy) else ((2,0) if dy>0 else (0,2)))
        straight=(sa,sb) in ((3,1),(1,3)) and abs(dy)<0.05 or (sa,sb) in ((2,0),(0,2)) and abs(dx)<0.05
        c=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT if straight else MSO_CONNECTOR.ELBOW,0,0,0,0)
        c.begin_connect(shp[a],sa); c.end_connect(shp[b],sb)
        c.line.color.rgb=RULE2; c.line.width=Pt(1.75)
        c.line._get_or_add_ln().append(c.line._get_or_add_ln().makeelement(qn("a:tailEnd"),{"type":"triangle","w":"med","len":"med"}))
        if label:  # centered above the first leg when it runs sideways, beside it when it runs down or up
            ax=A[0]+A[2]/2; ay=A[1]+A[3]/2
            if sa in (1,3):
                x0=A[0]+A[2] if sa==3 else A[0]
                x1=B[0] if straight and sa==3 else B[0]+B[2] if straight else x0+(0.9 if sa==3 else -0.9)
                text(s,(x0+x1)/2-0.9,ay-0.4,1.8,0.32,label,size=15,color=MUTE,align=CENTER,maxlines=1)
            else:
                y0=A[1]+A[3] if sa==2 else A[1]
                y1=B[1] if sb==0 else B[1]+B[3]/2
                text(s,ax+0.12,(y0+y1)/2-0.16,1.8,0.32,label,size=15,color=MUTE,maxlines=1)
    return shp

def pillars(s,cols,y=2.05,icon_dir="icons"):
    """Three to five parallel ideas, each under an icon in a pale disc: (icon name, headline, body).
    Make the icons with ../../scripts/icon.py NAME --circle E6E9FC --out icons."""
    n=len(cols); gap=0.45; cw=(CW-(n-1)*gap)/n; d=0.95
    for i,(icon,hd,body) in enumerate(cols):
        x=ML+i*(cw+gap)
        png=os.path.join(_D,icon_dir,icon+".png")
        if os.path.exists(png): s.shapes.add_picture(png,Inches(x),Inches(y),Inches(d),Inches(d))
        else: dot(s,x+d/2,y+d/2,d/2,PTINT)
        text(s,x,y+d+0.25,cw,0.5,hd,size=24,color=INK,bold=True,maxlines=1)
        rect(s,x,y+d+0.82,0.6,0.05,fill=PRIM)
        text(s,x,y+d+1.0,cw,2.0,body,size=20,color=INKM,sp=1.05)

# ============================ DEMO SLIDES (replace these) ============================
# A fictional case: Fernwood sells "home-compostable" coffee pods. The prompt: find the weakest
# point in the claim and say how the company should fix it. Everything below is invented.

def s_title():
    s=slide(NAVY); rect(s,0,0,W,0.16,fill=PRIM)
    for (r,lw,col) in [(2.9,2.0,RGBColor(0x2A,0x3B,0x7A)),(2.05,1.4,RGBColor(0x1E,0x2B,0x60)),(1.2,6.0,PRIM)]:
        ring(s,11.95,4.25,r,col,lw)
    text(s,ML,2.0,9.4,1.05,"Fernwood Pods",size=60,color=WHITE,bold=True,maxlines=1)
    text(s,ML,3.2,9.4,0.7,"Sold as compostable. Thrown away.",size=32,color=SKYL,maxlines=1)
    rect(s,ML,4.25,1.4,0.055,fill=PRIM)
    text(s,ML,4.55,9.4,0.5,"Team name  ·  Course or client  ·  Date",size=22,color=SKYD,maxlines=1)
    notes(s,"Open with the answer in one line: the subtitle. Then the prompt in your own words. About 30 seconds.")

def s_standard():
    s=slide(); head(s,"Standard of Comparison","How we rank the options")
    numbered_rows(s,[("Provable","the company’s own words and numbers prove it"),
                     ("Large","touches most of the product’s sales"),
                     ("Central","hits the claim the brand is built on"),
                     ("Hard to rebut","their best answer doesn’t close it")])
    banner(s,"Best answer = strong on all four tests",y=6.3,size=23)
    notes(s,"Before the answer, how we judge it. When two teams answer the same prompt, the one that sets the standard of comparison frames the debate. We return to this scorecard at the end. About 40 seconds.")

def s_argument():
    s=slide(); head(s,"The Gap","“Home-compostable” on a pod most buyers bin")
    text(s,ML,1.98,CW,0.9,"The compostable claim, on the 30-packs sold through grocery and office delivery",size=24,color=INK,bold=True,sp=1.02)
    label_rows(s,[("The box","“Home-compostable” in the largest type on the pack"),
                  ("The fine print","certified for industrial composting only"),
                  ("The channel","70% of sales = offices with no compost bin"),
                  ("The proof","no take-back and no composting rate published")])
    banner(s,"Four independent legs → break one and the rest still stand")
    src(s,"Fernwood packaging; annual report (illustrative)")
    notes(s,"The answer, and the legs it stands on. Each leg comes from a different source, so a rebuttal of one leaves the others. Name the source for each in the notes with page numbers; the slide shows only the source line. About 60 seconds.")

def s_table():
    s=slide(); head(s,"Pod vs Capsule","What the claim depends on")
    rule_table(s,["Benefit claimed","What it needs","Delivered?"],
        [("Breaks down in weeks","an industrial facility",("rarely",WARN)),
         ("Less landfill","the buyer to sort it",("rarely",WARN)),
         ("Fits existing machines","nothing",("yes",GOOD)),
         ("Premium feel","nothing",("yes",GOOD))],
        [3.6,5.0,CW-8.6])
    banner(s,"2 of 4 benefits need a step the buyer usually skips")
    src(s,"Product page; certification listing (illustrative)")

def s_scale():
    s=slide(); head(s,"Scale","Where the sales go")
    text(s,ML,1.95,CW,0.4,[("Sales by channel:  ",{"size":18,"color":MUTE,"bold":True}),
        ("70% go to offices without compost collection",{"size":18,"color":INK,"bold":True})],maxlines=1)
    share_bar(s,2.38,[(70,"office delivery",WARN,WHITE),(20,"grocery",GREY,INK),(10,"online",GREY,INK)])
    text(s,ML,3.62,CW,0.4,[("Sales by pack:  ",{"size":18,"color":MUTE,"bold":True}),
        ("the 30-pack is most of the volume",{"size":18,"color":INK,"bold":True})],maxlines=1)
    share_bar(s,4.05,[(60,"30-pack",PRIM,WHITE),(25,"12-pack",PRIMD,WHITE),(15,"sampler",GREY,INK)])
    big_numbers(s,[("$40M","annual pod sales"),("3","years since launch"),("0","published compost rate")])
    src(s,"Illustrative figures")

def s_journey():
    s=slide(); head(s,"Customer Decision Process","The claim sells before purchase and gets tested after")
    xs,cw=stages(s,[("Need","coffee at the office",None),("Search","the box on the shelf",None),
                    ("Evaluation","vs plain pods: the green claim wins here","accent"),
                    ("Purchase","a 30-pack on subscription",None),("Post-purchase","no bin, no instructions, into the trash","warn")])
    hline(s,xs[1],5.1,xs[3]+cw,PRIM,2.5)
    text(s,xs[1],5.2,xs[3]+cw-xs[1],0.4,"where the marketing budget goes",size=18,color=PRIM,bold=True,align=CENTER,maxlines=1)
    hline(s,xs[4],5.1,xs[4]+cw,WARN,2.5)
    text(s,xs[4],5.2,cw,0.62,"no budget here",size=18,color=WARN,bold=True,align=CENTER,maxlines=1)
    src(s,"Customer decision process (framework)")

def s_pillars():
    s=slide(); head(s,"The Fix","Three moves, one per leg")
    pillars(s,[("search","Audit the claim","every pack and page checked against the certification"),
               ("recycle","Collect the pods","mail-back bags for office plans, bins at partner offices"),
               ("badge-check","Prove it","an audited composting rate, published yearly")])
    banner(s,"Each move takes one leg away from the critic",y=6.2)
    src(s,"Icons: Lucide, drawn with scripts/icon.py")

def s_flow():
    s=slide(); head(s,"Where a Pod Goes","Today’s path, and the one the claim promises")
    flow(s,{"buy":(ML,2.3,2.3,0.9,"Office order",None),
            "use":(ML+3.25,2.3,2.3,0.9,"Used once","accent"),
            "bin":(ML+6.5,2.3,2.3,0.9,"Trash bin","warn"),
            "land":(ML+9.55,2.3,2.28,0.9,"Landfill","warn"),
            "mail":(ML+6.5,4.55,2.3,0.9,"Mail-back bag","ghost"),
            "comp":(ML+9.55,4.55,2.28,0.9,"Industrial compost","ghost")},
         [("buy","use",None),("use","bin","today"),("bin","land",None),("use","mail","proposed",(2,1)),("mail","comp",None)])
    banner(s,"Today every path ends in landfill → the fix adds the lower branch")

def s_score():
    s=slide(); head(s,"Every Other Option","Scored on the standard of comparison")
    scorecard(s,["Provable","Large","Central","Hard to rebut"],
        [("Ours: compostable claim × office channel",[3,3,3,3],True),
         ("Industrial-only fine print alone",[3,2,2,1],False),
         ("Packaging ink",[1,1,1,2],False),
         ("Celebrity endorsement",[2,1,1,1],False)])
    banner(s,"Each rival is one leg of ours → fixing it leaves ours standing",y=6.25)
    src(s,"●●● strong   ●●○ some   ●○○ weak   ·   the team’s scoring")

def s_assume():
    s=slide(); head(s,"Assumptions","What our case rests on, and the evidence")
    rule_table(s,["We assume","Evidence","If we’re wrong"],
        [("Offices don’t compost","no bin in the office channel survey","box and proof legs stand"),
         ("Buyers read the box, not the fine print","claim in large type, caveat in small","channel leg stands"),
         ("No take-back exists","none on the site or in the report","box leg stands")],
        [3.45,5.3,CW-8.75],rh=0.98,styles=[(INK,True),(INKM,False),(GOOD,True)])
    text(s,ML,6.4,CW,0.45,"No single assumption carries the argument",size=22,color=INK,bold=True,maxlines=1)

def s_rebut():
    s=slide(); head(s,"Rebuttals","Their best answers, and why the point holds")
    objection_rows(s,[("“The pods are certified compostable.”","certified for industrial sites, sold for homes"),
                      ("“Some buyers do compost them.”","70% of sales go where no one can"),
                      ("“It still beats plastic pods.”","the claim is “compostable”, judged on that word")])

def s_fix():
    s=slide(NAVY); rect(s,0,0,W,0.16,fill=PRIM)
    head(s,"Recommendation","Take each leg away from the critic",dark=True)
    horizons(s,[("Short term · now","Fix the words",["“industrial compost only” on the front","drop “home” from the claim"],
                 "removes the critic’s quote","marketing","claims signed off by legal"),
                ("Medium term · 6-12 months","Collect the pods",["mail-back bags for office plans","bins at partner offices"],
                 "makes the claim true where 70% sell","office sales","% of pods returned"),
                ("Long term · 12-24 months","Publish the number",["audited composting rate","checked by an outside group"],
                 "a verified number beats a slogan","sustainability lead","rate, published yearly")])
    src(s,"Recommendation is the team’s",dark=True)

def s_works():
    s=slide(); head(s,"Works Cited","Sources, readings and frameworks")
    for i,(lab,t) in enumerate([("Case","Author, A. “Case title.” Publisher, number, year."),
                                ("Reading","Author, B. “Article title.” Journal, issue, year."),
                                ("Framework","Three-level product; customer decision process.")]):
        cy=2.0+i*0.92
        text(s,ML,cy,2.3,0.92,lab,size=19,color=PRIM,bold=True,anchor=MID,maxlines=1)
        text(s,ML+2.45,cy,CW-2.45,0.92,t,size=19,color=INK,anchor=MID,sp=1.0)
        if i<2: hline(s,ML,cy+0.92,W-ML,RULE,0.75)
    src(s,"Page numbers are in the speaker notes.")

def s_backup():
    s=slide(); head(s,"Backup: Numbers","For questions")
    rule_table(s,["Figure","Value","Source"],[("Annual pod sales","$40M","report p. 3"),("Office share","70%","report p. 5")],
        [5.6,4.0,CW-9.6],rh=0.55,size=19,styles=[(INK,False),(PRIM,True),(MUTE,False)])
    notes(s,"Hidden backup for the question period: every number with its source page.")
    hide(s)

ORDER=[s_title,s_standard,s_argument,s_table,s_scale,s_journey,s_flow,s_score,s_assume,s_rebut,s_pillars,s_fix,s_works,s_backup]

# ---------------------------------------------------------------- save
def _repair(path):
    """python-pptx leaves the notesMaster unlinked; strict readers then call the file invalid."""
    zin=zipfile.ZipFile(path); pres=zin.read("ppt/presentation.xml").decode()
    rels=zin.read("ppt/_rels/presentation.xml.rels").decode()
    if "<p:notesMasterIdLst>" in pres: zin.close(); return False
    m=re.search(r'<Relationship Id="(rId\d+)" Type="[^"]*/notesMaster"', rels)
    if not m: zin.close(); return False
    ins='<p:notesMasterIdLst><p:notesMasterId r:id="%s"/></p:notesMasterIdLst>'%m.group(1)
    pres2=pres.replace("</p:sldMasterIdLst>","</p:sldMasterIdLst>"+ins,1)
    tmp=path+".tmp"; zo=zipfile.ZipFile(tmp,"w",zipfile.ZIP_DEFLATED)
    for it in zin.infolist():
        d=zin.read(it.filename)
        if it.filename=="ppt/presentation.xml": d=pres2.encode()
        zo.writestr(it,d)
    zo.close(); zin.close(); shutil.move(tmp,path); return True

if __name__=="__main__":
    pick=os.environ.get("SLIDES")
    idx=[int(v) for v in pick.split(",")] if pick else list(range(1,len(ORDER)+1))
    for n in idx:
        CUR["n"]=n; CUR["name"]=ORDER[n-1].__name__
        ORDER[n-1]()
    out=sys.argv[1] if len(sys.argv)>1 else os.path.join(_D,"deck.pptx")
    if SCALE!=1.0:
        for sl in prs.slides: _rescale(sl,SCALE)
    prs.save(out); fx=_repair(out)
    vis=sum(1 for sl in prs.slides if sl._element.get("show")!="0")
    print("saved",out,"; slides",len(prs.slides._sldIdLst),"visible",vis,"; notes relinked",fx)
    for w_ in WARNINGS: print("WARN",w_)
