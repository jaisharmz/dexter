#!/usr/bin/env python3
"""Read a slide template (.pptx or .potx) and report what a deck built on it must follow.

    uv run --no-project --with python-pptx python template.py company.potx
    ... template.py company.pptx --style STYLE.md     also write a STYLE.md whose tokens use the template's
                                                      fonts and colors, for the research and technical presets

A Google Slides template: File > Download > Microsoft PowerPoint (.pptx) first.

It prints the slide size, the theme's fonts and colors, every layout with its placeholders, and
the art on the master and layouts (logos, bars, footers) with the box each one covers, so content
can keep clear of it. The consulting preset builds on the template itself (TEMPLATE=company.potx
python deck.py out.pptx), so its masters, logos and footers carry over."""
import argparse
import os
import re
import shutil
import sys
import tempfile
import zipfile

from pptx import Presentation
from pptx.oxml.ns import qn

EMU = 914400


def open_template(path):
    """python-pptx refuses a .potx by its content type; a copy relabelled as a presentation opens."""
    with zipfile.ZipFile(path) as z:
        types = z.read("[Content_Types].xml").decode()
    if "presentationml.template.main+xml" not in types:
        return Presentation(path)
    tmp = tempfile.NamedTemporaryFile(suffix=".pptx", delete=False).name
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.decode().replace("presentationml.template.main+xml", "presentationml.presentation.main+xml").encode()
            zout.writestr(item, data)
    return Presentation(tmp)


def theme(prs):
    """Fonts and color scheme from the first master's theme part."""
    master = prs.slide_masters[0]
    part = next(r.target_part for r in master.part.rels.values() if r.reltype.endswith("/theme"))
    xml = part.blob.decode("utf8", "ignore")
    fonts = {k: (re.search(rf"<a:{k}Font>\s*<a:latin typeface=\"([^\"]*)\"", xml) or [None, None])[1] for k in ("major", "minor")}
    colors = {}
    scheme = re.search(r"<a:clrScheme[^>]*>(.*?)</a:clrScheme>", xml, re.S)
    for name in ("dk1", "lt1", "dk2", "lt2", "accent1", "accent2", "accent3", "accent4", "accent5", "accent6"):
        m = re.search(rf"<a:{name}>(.*?)</a:{name}>", scheme.group(1) if scheme else "", re.S)
        if m:
            v = re.search(r'(?:srgbClr val|lastClr)="([0-9A-Fa-f]{6})"', m.group(1))
            colors[name] = v.group(1).upper() if v else None
    return fonts, colors


def box(shape):
    try:
        return tuple(round((v or 0) / EMU, 2) for v in (shape.left, shape.top, shape.width, shape.height))
    except Exception:
        return None


def art(shapes):
    """Non-placeholder shapes: logos, bars, footer art."""
    out = []
    for sh in shapes:
        if sh.is_placeholder:
            continue
        kind = "picture" if sh.shape_type == 13 else "text" if sh.has_text_frame and sh.text_frame.text.strip() else "shape"
        out.append((kind, sh.name, box(sh), sh.text_frame.text.strip()[:40] if sh.has_text_frame else ""))
    return out


def blank_layout(prs):
    """The layout with the fewest placeholders (title-only beats one with a body); ties go to the name 'Blank'."""
    return min(prs.slide_layouts, key=lambda l: (len([p for p in l.placeholders if p.placeholder_format.type not in (13, 15, 16)]),
                                                 0 if "blank" in l.name.lower() else 1))


def lum(hex6):
    r, g, b = (int(hex6[i:i + 2], 16) / 255 for i in (0, 2, 4))
    return 0.2126 * r + 0.7152 * g + 0.0722 * b


def tokens(fonts, colors, w, h):
    """Preset tokens in the template's fonts and colors. background is empty so the master shows through."""
    body = fonts.get("minor") or "Arial"
    text = colors.get("dk1") or "000000"
    accent = colors.get("accent1") or "1155CC"
    muted = colors.get("dk2") if colors.get("dk2") and 0.15 < lum(colors["dk2"]) < 0.55 else "595959"
    return {
        "fonts": {"body": body, "bold": body, "title": fonts.get("major") or body, "code": "Courier New", "math": "STIX Two Math"},
        "sizes": {"title": 28, "subtitle": 20, "body": 18, "banner": 16, "caption": 14, "code": 14, "minimum": 11},
        "colors": {"text": text, "muted": muted, "accent": accent, "code": "38761D", "codeHighlight": accent,
                   "box": "F2F2F2" if (colors.get("lt2") or "FFFFFF") == "FFFFFF" else colors["lt2"], "pale": "E8EEF8",
                   "good": "38761D", "bad": "980000", "imageBorder": "B7B7B7", "arrow": "595959", "background": ""},
        "banner": {"fill": accent, "text": "FFFFFF" if lum(accent) < 0.6 else "000000"},
        "layout": {"titleX": 0.43, "titleY": 0.3, "margin": 0.5, "bannerY": 4.74, "bannerW": 8.7, "bannerH": 0.56},
        "bullets": ["25CF", "25CB", "2605"],
        "flags": {"agent": ["Agent", "FCE5CD", "B45F06"], "script": ["Script", "D9EAD3", "38761D"], "input": ["Input", "EFEFEF", "595959"]},
        "template": {"width": w, "height": h},
    }


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("template")
    ap.add_argument("--style", help="write a STYLE.md with the template's tokens")
    a = ap.parse_args()
    prs = open_template(a.template)
    w, h = round(prs.slide_width / EMU, 3), round(prs.slide_height / EMU, 3)
    fonts, colors = theme(prs)
    print(f"size      {w} x {h} in ({'16:9' if abs(w / h - 16 / 9) < 0.01 else '4:3' if abs(w / h - 4 / 3) < 0.01 else 'other'})")
    print(f"fonts     headings {fonts.get('major')}, body {fonts.get('minor')}")
    print("colors    " + "  ".join(f"{k} {v}" for k, v in colors.items()))
    print("layouts")
    for i, lay in enumerate(prs.slide_layouts):
        ph = ", ".join(f"{p.placeholder_format.type}".split(" ")[0].lower() + f"@{box(p)}" for p in lay.placeholders)
        print(f"  {i:2d} {lay.name!r}: {ph or 'no placeholders'}")
    print(f"blank     {blank_layout(prs).name!r} (what the builders add slides on)")
    print("master art (keep content clear of these boxes: x, y, w, h in inches)")
    for kind, name, b, t in art(prs.slide_masters[0].shapes) + art(blank_layout(prs).shapes):
        print(f"  {kind:8s} {name!r} {b} {t!r}")
    print(f"existing  {len(prs.slides)} slides (the consulting builder removes them)")
    if a.style:
        import json
        body = json.dumps(tokens(fonts, colors, w, h), indent=2)
        with open(a.style, "w") as f:
            f.write(f"# Template style\n\nTokens read from `{os.path.basename(a.template)}` by scripts/template.py. "
                    f"Fonts and colors are the template's; edit the sizes if its body text runs larger.\n\n```json\n{body}\n```\n")
        print("wrote", a.style)


if __name__ == "__main__":
    main()
